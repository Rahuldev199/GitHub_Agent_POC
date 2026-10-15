from __future__ import annotations

import io
from datetime import UTC, datetime

from openpyxl import load_workbook

from github_team_stats.domain.activity import ActivityTotals
from github_team_stats.domain.team import TeamActivityReport, TeamMemberActivity
from github_team_stats.presentation.web.exports import csv_export, html_export, xlsx_export


def report(login: str = "alice") -> TeamActivityReport:
    now = datetime.now(UTC)
    totals = ActivityTotals(
        commits=0, lines_added=0, lines_deleted=0, lines_changed=0,
        pull_requests_opened=0, pull_requests_merged=0, reviews_submitted=0, workflow_runs=0,
    )
    return TeamActivityReport(
        team_key="backend", team_name="Backend", requester_login="lead", period="2026-10",
        scope="repos:acme/api", members_requested=1, members_completed=0, totals=totals,
        members=[TeamMemberActivity(member_login=login, cache_status="failed")], complete=False,
        warnings=[], generated_at=now, data_as_of=now,
    )


def test_exports_are_self_contained() -> None:
    value = report()
    assert b"Member,Commits" in csv_export(value)
    page = html_export(value)
    assert b"<!doctype html>" in page
    assert b"http://" not in page and b"https://" not in page
    workbook = load_workbook(io.BytesIO(xlsx_export(value)))
    assert workbook["Members"]["A2"].value == "alice"


def test_xlsx_blocks_formula_injection() -> None:
    malicious = report("=2+2")
    workbook = load_workbook(io.BytesIO(xlsx_export(malicious)))
    assert workbook["Members"]["A2"].value == "'=2+2"
    assert b"'=2+2" in csv_export(malicious)
