"""Safe, in-memory team report exports."""

from __future__ import annotations

import csv
import html
import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from github_team_stats.domain.team import TeamActivityReport

HEADERS = [
    "Member", "Commits", "Lines added", "Lines deleted", "Lines changed",
    "PRs opened", "PRs merged", "Reviews", "Workflow runs", "Complete",
]
FIELDS = [
    "commits", "lines_added", "lines_deleted", "lines_changed", "pull_requests_opened",
    "pull_requests_merged", "reviews_submitted", "workflow_runs",
]


def report_rows(report: TeamActivityReport) -> list[list[object]]:
    rows: list[list[object]] = []
    for member in report.members:
        totals = member.activity.totals if member.activity else None
        rows.append(
            [member.member_login, *(getattr(totals, field, 0) if totals else 0 for field in FIELDS),
             bool(member.activity and member.activity.complete)]
        )
    return rows


def csv_export(report: TeamActivityReport) -> bytes:
    output = io.StringIO(newline="")
    writer = csv.writer(output, lineterminator="\n")
    writer.writerow(HEADERS)
    writer.writerows([_safe_cell(value) for value in row] for row in report_rows(report))
    return output.getvalue().encode("utf-8-sig")


def xlsx_export(report: TeamActivityReport) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    if not isinstance(sheet, Worksheet):
        raise RuntimeError("Workbook did not create a worksheet")
    sheet.title = "Members"
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="075E66")
    for row in report_rows(report):
        sheet.append([_safe_cell(value) for value in row])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, column in enumerate(sheet.columns, start=1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = min(
            28, max(12, max(len(str(cell.value or "")) for cell in column) + 2)
        )
    warnings = workbook.create_sheet("Warnings")
    warnings.append(["Member", "Repository", "Source", "Code", "Message"])
    for warning in report.warnings:
        warnings.append(
            [_safe_cell(warning.member_login), _safe_cell(warning.repository), warning.source,
             warning.code, _safe_cell(warning.message)]
        )
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def html_export(report: TeamActivityReport) -> bytes:
    rows = "".join(
        "<tr>" + "".join(f"<td>{html.escape(str(value))}</td>" for value in row) + "</tr>"
        for row in report_rows(report)
    )
    warning_rows = "".join(
        f"<li><b>{html.escape(w.member_login)}</b> — {html.escape(w.repository)} / "
        f"{html.escape(w.source)} ({w.code}): {html.escape(w.message)}</li>"
        for w in report.warnings
    ) or "<li>None</li>"
    headings = "".join(f"<th>{html.escape(header)}</th>" for header in HEADERS)
    status = "Complete" if report.complete else "Incomplete"
    document = "".join(
        [
            '<!doctype html><html><head><meta charset="utf-8"><title>Team activity</title>',
            "<style>body{font:15px system-ui;margin:40px;color:#16252a}",
            "h1{color:#075e66}table{border-collapse:collapse;width:100%}",
            "th,td{padding:9px;border-bottom:1px solid #d8e3e5;text-align:right}",
            "th:first-child,td:first-child{text-align:left}",
            "th{background:#075e66;color:white}</style></head><body>",
            f"<h1>{html.escape(report.team_name)} activity</h1>",
            f"<p>Period: {html.escape(report.period)} · Read-only snapshot · {status}</p>",
            f"<table><thead><tr>{headings}</tr></thead><tbody>{rows}</tbody></table>",
            f"<h2>Warnings</h2><ul>{warning_rows}</ul></body></html>",
        ]
    )
    return document.encode("utf-8")


def _safe_cell(value: object) -> object:
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value
